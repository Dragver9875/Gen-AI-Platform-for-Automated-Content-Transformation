from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from artifacts.models import PDFArtifactIR, PresentationIR, SVGArtifactIR, TextArtifactIR


def _safe_slug(text: str) -> str:
    cleaned = "".join(ch.lower() if ch.isalnum() else "-" for ch in text).strip("-")
    while "--" in cleaned:
        cleaned = cleaned.replace("--", "-")
    return cleaned[:72] or "artifact"


class TextSerializer:
    def serialize(self, ir: TextArtifactIR, destination: Path) -> Path:
        destination.write_text(ir.content, encoding="utf-8")
        return destination


def _validate_generated_typst(source: str) -> None:
    lowered = source.lower()
    forbidden = ("#import", "#include", "read(", "csv(", "json(", "yaml(", "xml(", "http://", "https://")
    hit = next((token for token in forbidden if token in lowered), None)
    if hit:
        raise ValueError(f"Generated Typst contains forbidden external-resource primitive: {hit}")


def _validate_generated_svg(source: str) -> None:
    lowered = source.lower()
    forbidden = ("<script", "<foreignobject", "javascript:", "xlink:href", 'href="http://', 'href="https://', "href='http://", "href='https://")
    hit = next((token for token in forbidden if token in lowered), None)
    if hit:
        raise ValueError(f"Generated SVG contains forbidden active/external content: {hit}")


class TypstPDFSerializer:
    def __init__(self, *, typst_binary: str = "typst", retain_source: bool = True):
        self.typst_binary = typst_binary
        self.retain_source = retain_source

    def serialize(self, ir: PDFArtifactIR, destination: Path) -> tuple[Path | None, Path]:
        _validate_generated_typst(ir.typst_source)
        source_path = destination.with_suffix(".typ")
        source_path.write_text(ir.typst_source, encoding="utf-8")
        binary = shutil.which(self.typst_binary) or (self.typst_binary if Path(self.typst_binary).exists() else None)
        if not binary:
            return None, source_path
        subprocess.run([binary, "compile", str(source_path), str(destination)], check=True, capture_output=True, text=True)
        if not self.retain_source:
            source_path.unlink(missing_ok=True)
        return destination, source_path


class PPTXSerializer:
    def serialize(self, ir: PresentationIR, destination: Path) -> Path:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("python-pptx is required to serialize PPTX artifacts") from exc

        prs = Presentation()
        # Use conservative built-in layouts; the ML-generated layout remains
        # metadata/intent while serialization stays deterministic and editable.
        for index, slide_ir in enumerate(ir.slides):
            layout_index = 0 if index == 0 and not slide_ir.bullets and not slide_ir.body else 1
            slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
            if slide.shapes.title:
                slide.shapes.title.text = slide_ir.title
            if layout_index == 0:
                subtitle = slide.placeholders[1] if len(slide.placeholders) > 1 else None
                if subtitle is not None:
                    subtitle.text = slide_ir.subtitle or slide_ir.body
            else:
                body_shape = slide.placeholders[1] if len(slide.placeholders) > 1 else None
                if body_shape is not None:
                    frame = body_shape.text_frame
                    frame.clear()
                    if slide_ir.body:
                        frame.text = slide_ir.body
                    for bullet in slide_ir.bullets:
                        p = frame.add_paragraph()
                        p.text = bullet
                        p.level = 0
                    for metric in slide_ir.metrics:
                        p = frame.add_paragraph()
                        p.text = f"{metric.get('value', '')} — {metric.get('label', '')}".strip(" —")
                        p.level = 0
            if slide_ir.speaker_notes:
                try:
                    notes = slide.notes_slide.notes_text_frame
                    notes.text = slide_ir.speaker_notes
                except Exception:
                    pass
        prs.save(destination)
        return destination


class SVGSerializer:
    def serialize(self, ir: SVGArtifactIR, destination: Path) -> Path:
        _validate_generated_svg(ir.svg)
        destination.write_text(ir.svg, encoding="utf-8")
        return destination
