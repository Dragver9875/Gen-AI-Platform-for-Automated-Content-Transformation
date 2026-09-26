from __future__ import annotations

from pathlib import Path

import pymupdf
from pptx import Presentation

from ingestion.chunker import StructureAwareChunker
from ingestion.router import IngestionRouter


class FakeVLM:
    model = "Qwen/Qwen2.5-VL-3B-Instruct"

    def __init__(self):
        self.byte_calls = 0
        self.file_calls = 0
        self.last_model_used = self.model
        self.prompts: list[str] = []

    def describe_file(self, path, prompt=None):
        self.file_calls += 1
        self.prompts.append(prompt or "")
        return "Visible image content with label 42."

    def describe_bytes(self, data, prompt=None, *, media_type=None):
        self.byte_calls += 1
        self.prompts.append(prompt or "")
        return "# Encoded page\n\nProject finished today. Metric: 42."


def make_router(*, vlm=None, max_pdf_pages=40, max_pptx_slides=40):
    return IngestionRouter(
        vlm=vlm or FakeVLM(),
        render_dpi=96,
        max_pdf_pages=max_pdf_pages,
        max_pptx_slides=max_pptx_slides,
    )


def test_text_ingestion_stays_direct(tmp_path: Path):
    file = tmp_path / "source.txt"
    file.write_text("Alpha beta gamma", encoding="utf-8")
    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(file)
    assert result.strategy == "text-direct"
    assert result.text == "Alpha beta gamma"
    assert vlm.byte_calls == 0
    assert vlm.file_calls == 0


def test_text_chunking_preserves_tenant_metadata(tmp_path: Path):
    file = tmp_path / "source.txt"
    file.write_text("Alpha beta gamma", encoding="utf-8")
    result = make_router().ingest(file)
    chunks = StructureAwareChunker(target_chars=100, overlap_chars=0).chunk(
        result, user_id="u1", session_id="s1"
    )
    assert len(chunks) == 1
    assert chunks[0].metadata["user_id"] == "u1"
    assert chunks[0].metadata["session_id"] == "s1"


def test_every_image_goes_directly_through_shared_qwen_encoder(tmp_path: Path):
    file = tmp_path / "photo.jpeg"
    file.write_bytes(b"fake-image-bytes")
    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(file)
    assert result.strategy == "qwen2.5-vl-image"
    assert vlm.file_calls == 1
    assert result.provider_metadata["shared_encoder"] == vlm.model
    assert "label 42" in result.text


def test_document_page_image_uses_same_shared_encoder_no_router(tmp_path: Path):
    file = tmp_path / "pdf2photo.jpeg"
    file.write_bytes(b"fake-document-image")
    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(file)
    assert result.strategy == "qwen2.5-vl-image"
    assert vlm.file_calls == 1
    assert "scanned document page" in vlm.prompts[0]


def test_native_pdf_also_goes_through_qwen_per_page(tmp_path: Path):
    pdf_path = tmp_path / "native.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Native searchable PDF text.")
    doc.save(pdf_path)
    doc.close()

    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(pdf_path)
    assert result.strategy == "qwen2.5-vl-pdf"
    assert vlm.byte_calls == 1
    assert result.provider_metadata["pages_encoded"] == 1
    assert result.provider_metadata["shared_encoder"] == vlm.model
    assert "Project finished today" in result.text


def test_multipage_pdf_calls_qwen_for_each_page(tmp_path: Path):
    pdf_path = tmp_path / "multi.pdf"
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i+1}")
    doc.save(pdf_path)
    doc.close()

    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(pdf_path)
    assert vlm.byte_calls == 3
    assert len(result.elements) == 3
    assert [e.page for e in result.elements] == [1, 2, 3]


def test_pdf_page_cap_is_visible_not_silent(tmp_path: Path):
    pdf_path = tmp_path / "long.pdf"
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page().insert_text((72, 72), f"Page {i+1}")
    doc.save(pdf_path)
    doc.close()

    vlm = FakeVLM()
    result = make_router(vlm=vlm, max_pdf_pages=2).ingest(pdf_path)
    assert vlm.byte_calls == 2
    assert result.provider_metadata["total_pages"] == 3
    assert any("capped at 2" in warning for warning in result.warnings)


def test_every_pptx_slide_is_rasterized_then_encoded_by_qwen(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    for title in ("Quarterly Update", "Next Steps"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Project finished today."
    prs.save(path)

    vlm = FakeVLM()
    result = make_router(vlm=vlm).ingest(path)
    assert result.strategy == "qwen2.5-vl-pptx"
    assert vlm.byte_calls == 2
    assert result.provider_metadata["slides_encoded"] == 2
    assert result.provider_metadata["slide_renderer"] == "python-pptx+pillow"
    assert [e.slide for e in result.elements] == [1, 2]


def test_pptx_slide_cap_is_visible(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i+1}"
        slide.placeholders[1].text = "Body"
    prs.save(path)

    vlm = FakeVLM()
    result = make_router(vlm=vlm, max_pptx_slides=2).ingest(path)
    assert vlm.byte_calls == 2
    assert result.provider_metadata["total_slides"] == 3
    assert any("capped at 2" in warning for warning in result.warnings)
