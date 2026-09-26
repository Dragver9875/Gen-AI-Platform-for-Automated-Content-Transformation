from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ingestion.normalizer import normalize_text


def _font(size: int) -> ImageFont.ImageFont:
    candidates = [
        r"C:\\Windows\\Fonts\\arial.ttf",
        r"C:\\Windows\\Fonts\\calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        try:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size=size)
        except Exception:
            pass
    return ImageFont.load_default()


def render_pdf_pages(path: str | Path, *, dpi: int = 144, max_pages: int = 0) -> list[tuple[int, bytes]]:
    """Render PDF pages to PNG bytes for the shared multimodal encoder.

    max_pages=0 means unlimited. The caller is responsible for warning when a
    user-selected cap truncates a document.
    """
    document = pymupdf.open(path)
    pages: list[tuple[int, bytes]] = []
    try:
        limit = document.page_count if max_pages <= 0 else min(document.page_count, max_pages)
        scale = max(72, int(dpi)) / 72.0
        matrix = pymupdf.Matrix(scale, scale)
        for index in range(limit):
            page = document.load_page(index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            pages.append((index + 1, pix.tobytes("png")))
    finally:
        document.close()
    return pages


def _chart_text(shape: Any) -> str:
    try:
        chart = shape.chart
        lines: list[str] = []
        if getattr(chart, "has_title", False):
            title = normalize_text(chart.chart_title.text_frame.text or "")
            if title:
                lines.append(f"Chart: {title}")
        for plot in chart.plots:
            try:
                categories = [str(category) for category in plot.categories]
            except Exception:
                categories = []
            for series in plot.series:
                name = str(getattr(series, "name", "Series") or "Series")
                values = list(getattr(series, "values", []) or [])
                if categories and values:
                    lines.append(name + " — " + "; ".join(f"{c}: {v}" for c, v in zip(categories, values)))
                elif values:
                    lines.append(name + " — " + ", ".join(str(value) for value in values))
        return normalize_text("\n".join(lines))
    except Exception:
        return ""


def _fit_box(shape: Any, slide_width: int, slide_height: int, canvas_width: int, canvas_height: int) -> tuple[int, int, int, int]:
    x = int((int(getattr(shape, "left", 0)) / slide_width) * canvas_width)
    y = int((int(getattr(shape, "top", 0)) / slide_height) * canvas_height)
    w = max(1, int((int(getattr(shape, "width", 1)) / slide_width) * canvas_width))
    h = max(1, int((int(getattr(shape, "height", 1)) / slide_height) * canvas_height))
    return x, y, min(canvas_width - x, w), min(canvas_height - y, h)


def _draw_wrapped(draw: ImageDraw.ImageDraw, text: str, box: tuple[int, int, int, int], *, font: ImageFont.ImageFont) -> None:
    x, y, w, h = box
    if not text or w <= 4 or h <= 4:
        return
    avg_char = max(5, int(getattr(font, "size", 14) * 0.55))
    max_chars = max(8, int(w / avg_char))
    words = text.replace("\n", " \n ").split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        if word == "\n":
            if current:
                lines.append(" ".join(current))
                current = []
            continue
        candidate = " ".join([*current, word])
        if len(candidate) > max_chars and current:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    line_height = max(13, int(getattr(font, "size", 14) * 1.25))
    max_lines = max(1, h // line_height)
    clipped = lines[:max_lines]
    if len(lines) > max_lines and clipped:
        clipped[-1] = clipped[-1][: max(0, max_chars - 1)] + "…"
    draw.multiline_text((x + 3, y + 3), "\n".join(clipped), fill="black", font=font, spacing=2)


def render_pptx_slides(
    path: str | Path,
    *,
    width: int = 1280,
    height: int = 720,
    max_slides: int = 0,
) -> tuple[list[tuple[int, bytes]], int]:
    """Create dependency-light slide rasterizations from PPTX contents.

    This is intentionally approximate rather than a PowerPoint renderer. It
    preserves text-box positions, tables, chart data, and embedded pictures so
    every slide can still pass through the same Qwen2.5-VL encoder without
    requiring PowerPoint/LibreOffice on the host.
    """
    presentation = Presentation(str(path))
    total = len(presentation.slides)
    limit = total if max_slides <= 0 else min(total, max_slides)
    slide_width = max(1, int(presentation.slide_width))
    slide_height = max(1, int(presentation.slide_height))
    rendered: list[tuple[int, bytes]] = []

    for slide_index, slide in enumerate((presentation.slides[i] for i in range(limit)), start=1):
        canvas = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(canvas)
        draw.rectangle((0, 0, width - 1, height - 1), outline="#d0d0d0", width=2)

        for shape in slide.shapes:
            x, y, w, h = _fit_box(shape, slide_width, slide_height, width, height)
            if w <= 0 or h <= 0:
                continue

            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
                    image.thumbnail((w, h))
                    px = x + max(0, (w - image.width) // 2)
                    py = y + max(0, (h - image.height) // 2)
                    canvas.paste(image, (px, py))
                except Exception:
                    draw.rectangle((x, y, x + w, y + h), outline="#888888", width=1)
                    _draw_wrapped(draw, "[Embedded picture]", (x, y, w, h), font=_font(14))
                continue

            if getattr(shape, "has_table", False):
                try:
                    rows = [[normalize_text(cell.text) for cell in row.cells] for row in shape.table.rows]
                    table_text = "\n".join(" | ".join(row) for row in rows if any(row))
                    draw.rectangle((x, y, x + w, y + h), outline="#777777", width=1)
                    _draw_wrapped(draw, table_text, (x, y, w, h), font=_font(14))
                except Exception:
                    pass
                continue

            if getattr(shape, "has_chart", False):
                chart_text = _chart_text(shape)
                draw.rectangle((x, y, x + w, y + h), outline="#777777", width=1)
                _draw_wrapped(draw, chart_text or "[Chart]", (x, y, w, h), font=_font(14))
                continue

            if getattr(shape, "has_text_frame", False):
                text = normalize_text(getattr(shape, "text", "") or "")
                if text:
                    estimated = max(14, min(30, int(h / max(2, text.count("\n") + 2))))
                    _draw_wrapped(draw, text, (x, y, w, h), font=_font(estimated))

        output = io.BytesIO()
        canvas.save(output, format="PNG", optimize=True)
        rendered.append((slide_index, output.getvalue()))

    return rendered, total
