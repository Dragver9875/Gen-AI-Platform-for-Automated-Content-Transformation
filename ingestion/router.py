from __future__ import annotations

import hashlib
import mimetypes
from pathlib import Path
from typing import Any

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas import IngestionResult, SourceElement
from ingestion.normalizer import normalize_text
from ingestion.image_heuristics import assess_document_image
from ingestion.preflight import PdfPreflight
from ingestion.registry import IngestionProcessorRegistry
from providers.siglip import SigLIPRoutingProvider
from providers.vlm import VLMProvider


_TEXT_SUFFIXES = {".txt", ".md"}
_DOCUMENT_SUFFIXES = {".pdf", ".pptx"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
DEFAULT_VISUAL_LABELS = ("document page", "screenshot", "chart", "diagram", "map", "photograph", "other visual")


class IngestionRouter:
    """Adaptive multimodal ingestion without a mandatory document-model endpoint.

    Design:
      * TXT/MD: direct deterministic extraction.
      * Native PDF pages: PyMuPDF text extraction (zero model calls).
      * Scanned/visual/mixed PDF pages: render only the pages that need vision and
        send those page images to the configured HF-compatible VLM.
      * PPTX: deterministic text/table/chart extraction with python-pptx; embedded
        pictures are described by the VLM.
      * Images: a deterministic document-page precheck runs first. SigLIP is optional;
        when disabled/unavailable, the configured VLM performs the lightweight routing.

    This deliberately avoids Granite Docling as a serverless default because that
    checkpoint may exist on the Hub without being deployed by an Inference Provider.
    """

    def __init__(
        self,
        *,
        siglip: SigLIPRoutingProvider | None,
        vlm: VLMProvider,
        pdf_preflight: PdfPreflight,
        enable_pdf_visual_fallback: bool = True,
        pdf_visual_fallback_max_pages: int = 12,
        document_like_image_labels: set[str] | None = None,
        pdf_visual_prompt: str | None = None,
    ):
        self.siglip = siglip
        self.vlm = vlm
        self.pdf_preflight = pdf_preflight
        self.enable_pdf_visual_fallback = enable_pdf_visual_fallback
        self.pdf_visual_fallback_max_pages = max(0, int(pdf_visual_fallback_max_pages))
        self.document_like_image_labels = {
            x.strip().lower()
            for x in (document_like_image_labels or {"document page", "screenshot"})
            if x.strip()
        }
        self.pdf_visual_prompt = pdf_visual_prompt or (
            "Describe this document page faithfully for retrieval. Preserve visible text, labels, "
            "numbers, chart trends, diagram relationships, and important visual content. "
            "If it is primarily a scanned text page, transcribe the meaningful content."
        )
        self.processors = IngestionProcessorRegistry()
        self.processors.register("text", _TEXT_SUFFIXES, self._ingest_text)
        self.processors.register("document", _DOCUMENT_SUFFIXES, self._ingest_document)
        self.processors.register("image", _IMAGE_SUFFIXES, self._ingest_image)

    def register_processor(self, name: str, suffixes: set[str], handler) -> None:
        self.processors.register(name, suffixes, handler)

    def ingest(self, path: str | Path) -> IngestionResult:
        file_path = Path(path)
        if not file_path.is_file():
            raise FileNotFoundError(file_path)
        source_id = hashlib.sha1(file_path.read_bytes()).hexdigest()[:20]
        media_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
        processor = self.processors.resolve(file_path)
        if processor is None:
            raise ValueError(
                f"Unsupported input type: {file_path.suffix.lower() or media_type}. "
                f"Supported suffixes: {', '.join(self.processors.supported_suffixes())}"
            )
        return processor.handler(file_path, source_id, media_type)

    def _ingest_text(self, file_path: Path, source_id: str, media_type: str) -> IngestionResult:
        raw = file_path.read_text(encoding="utf-8", errors="replace")
        text = normalize_text(raw)
        return IngestionResult(
            source_id,
            file_path.name,
            media_type,
            "text-direct",
            [SourceElement("text-0", "paragraph", text, raw_text=raw)],
        )

    def _ingest_document(self, path: Path, source_id: str, media_type: str) -> IngestionResult:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._ingest_pdf(path, source_id, media_type)
        if suffix == ".pptx":
            return self._ingest_pptx(path, source_id, media_type)
        raise ValueError(f"Unsupported document input: {suffix}")

    def _ingest_pdf(self, path: Path, source_id: str, media_type: str) -> IngestionResult:
        profile = self.pdf_preflight.inspect(path)
        page_profiles = {p.page: p for p in profile.pages}
        elements: list[SourceElement] = []
        warnings: list[str] = []
        native_pages = 0
        vlm_pages = 0
        vlm_models_used: set[str] = set()

        doc = pymupdf.open(path)
        try:
            for page_number, page in enumerate(doc, start=1):
                raw_text = page.get_text("text") or ""
                text = normalize_text(raw_text)
                if text:
                    native_pages += 1
                    elements.append(
                        SourceElement(
                            element_id=f"pdf-text-{page_number}",
                            kind="page_text",
                            text=text,
                            raw_text=raw_text,
                            page=page_number,
                            metadata={"processor": "pymupdf"},
                        )
                    )
        finally:
            doc.close()

        visual_candidates = profile.visual_or_scanned_pages
        selected_visual_pages = visual_candidates[: self.pdf_visual_fallback_max_pages]
        if len(visual_candidates) > len(selected_visual_pages):
            warnings.append(
                f"PDF has {len(visual_candidates)} visual/scanned page(s); VLM analysis was capped at "
                f"{len(selected_visual_pages)} by PDF_VISUAL_FALLBACK_MAX_PAGES."
            )

        if self.enable_pdf_visual_fallback:
            for page_number in selected_visual_pages:
                page_profile = page_profiles.get(page_number)
                kind = page_profile.kind if page_profile else "visual"
                try:
                    png = self.pdf_preflight.render_page_png(path, page_number)
                    if kind in {"scanned", "visual"}:
                        prompt = self.pdf_visual_prompt
                    else:
                        prompt = (
                            "Describe the non-text visual information on this mixed document page for retrieval: "
                            "charts, figures, diagrams, maps, relationships, labels and values. Avoid repeating "
                            "ordinary body text unless it is necessary to understand the visual."
                        )
                    description = normalize_text(self.vlm.describe_bytes(png, prompt=prompt, media_type="image/png"))
                    if description:
                        vlm_pages += 1
                        if getattr(self.vlm, "last_model_used", None):
                            vlm_models_used.add(str(self.vlm.last_model_used))
                        elements.append(
                            SourceElement(
                                element_id=f"pdf-visual-{page_number}",
                                kind="page_visual" if kind == "mixed" else "page_ocr_visual",
                                text=description,
                                page=page_number,
                                metadata={"processor": "vlm", "page_kind": kind},
                            )
                        )
                except Exception as exc:
                    # Native text remains useful for mixed/native documents. For a purely
                    # scanned PDF, surface a strong warning instead of aborting the whole file.
                    warnings.append(f"VLM page analysis failed on page {page_number}: {exc}")

        if not elements:
            raise RuntimeError(
                "PDF contained no extractable native text and no VLM page analysis succeeded. "
                "Check VLM availability or increase PDF_VISUAL_FALLBACK_MAX_PAGES."
            )

        return IngestionResult(
            source_id,
            path.name,
            media_type,
            f"pdf-{profile.strategy}",
            elements,
            warnings=warnings,
            provider_metadata={
                "page_profiles": [p.__dict__ for p in profile.pages],
                "native_text_pages": native_pages,
                "vlm_pages": vlm_pages,
                "vlm_models_used": sorted(vlm_models_used),
                "document_parser": "pymupdf+vlm",
            },
        )

    def _ingest_pptx(self, path: Path, source_id: str, media_type: str) -> IngestionResult:
        presentation = Presentation(str(path))
        elements: list[SourceElement] = []
        warnings: list[str] = []
        picture_count = 0
        picture_vlm_success = 0
        chart_count = 0
        vlm_models_used: set[str] = set()

        for slide_number, slide in enumerate(presentation.slides, start=1):
            title_shape = getattr(slide.shapes, "title", None)
            title_text = normalize_text(getattr(title_shape, "text", "") or "") if title_shape is not None else ""
            if title_text:
                elements.append(
                    SourceElement(
                        f"pptx-title-{slide_number}",
                        "heading",
                        title_text,
                        slide=slide_number,
                        section=title_text,
                        metadata={"processor": "python-pptx"},
                    )
                )

            for shape_index, shape in enumerate(slide.shapes):
                if shape is title_shape:
                    continue

                if getattr(shape, "has_text_frame", False):
                    raw = getattr(shape, "text", "") or ""
                    text = normalize_text(raw)
                    if text:
                        elements.append(
                            SourceElement(
                                f"pptx-text-{slide_number}-{shape_index}",
                                "slide_text",
                                text,
                                raw_text=raw,
                                slide=slide_number,
                                section=title_text or None,
                                metadata={"processor": "python-pptx"},
                            )
                        )

                if getattr(shape, "has_table", False):
                    try:
                        rows = [[normalize_text(cell.text) for cell in row.cells] for row in shape.table.rows]
                        table_text = "\n".join(" | ".join(row) for row in rows if any(row)).strip()
                        if table_text:
                            elements.append(
                                SourceElement(
                                    f"pptx-table-{slide_number}-{shape_index}",
                                    "table",
                                    table_text,
                                    slide=slide_number,
                                    section=title_text or None,
                                    metadata={"processor": "python-pptx"},
                                )
                            )
                    except Exception as exc:
                        warnings.append(f"Could not extract table on slide {slide_number}: {exc}")

                if getattr(shape, "has_chart", False):
                    chart_count += 1
                    chart_text = self._chart_to_text(shape)
                    if chart_text:
                        elements.append(
                            SourceElement(
                                f"pptx-chart-{slide_number}-{shape_index}",
                                "chart_data",
                                chart_text,
                                slide=slide_number,
                                section=title_text or None,
                                metadata={"processor": "python-pptx"},
                            )
                        )

                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                    try:
                        image = shape.image
                        description = normalize_text(
                            self.vlm.describe_bytes(
                                image.blob,
                                media_type=image.content_type,
                                prompt=(
                                    "Describe this presentation image faithfully for retrieval. Preserve visible text, "
                                    "numbers, labels, chart/diagram relationships and relevant visual context. "
                                    "Do not invent details."
                                ),
                            )
                        )
                        if description:
                            picture_vlm_success += 1
                            if getattr(self.vlm, "last_model_used", None):
                                vlm_models_used.add(str(self.vlm.last_model_used))
                            elements.append(
                                SourceElement(
                                    f"pptx-picture-{slide_number}-{shape_index}",
                                    "picture_description",
                                    description,
                                    slide=slide_number,
                                    section=title_text or None,
                                    metadata={"processor": "vlm", "content_type": image.content_type},
                                )
                            )
                    except Exception as exc:
                        warnings.append(f"VLM analysis failed for picture on slide {slide_number}: {exc}")

        if not elements:
            raise RuntimeError("PPTX contained no extractable text, tables, chart data, or VLM-readable pictures.")

        return IngestionResult(
            source_id,
            path.name,
            media_type,
            "pptx-native+visual",
            elements,
            warnings=warnings,
            provider_metadata={
                "document_parser": "python-pptx+vlm",
                "slides": len(presentation.slides),
                "embedded_pictures": picture_count,
                "pictures_analyzed_by_vlm": picture_vlm_success,
                "vlm_models_used": sorted(vlm_models_used),
                "charts": chart_count,
            },
        )

    @staticmethod
    def _chart_to_text(shape: Any) -> str:
        try:
            chart = shape.chart
            lines: list[str] = []
            if getattr(chart, "chart_title", None) is not None and getattr(chart, "has_title", False):
                title = normalize_text(chart.chart_title.text_frame.text or "")
                if title:
                    lines.append(f"Chart: {title}")
            for plot in chart.plots:
                categories = []
                try:
                    categories = [str(category) for category in plot.categories]
                except Exception:
                    categories = []
                for series in plot.series:
                    name = str(getattr(series, "name", "Series") or "Series")
                    values = list(getattr(series, "values", []) or [])
                    if categories and values:
                        pairs = [f"{c}: {v}" for c, v in zip(categories, values)]
                        lines.append(f"{name} — " + "; ".join(pairs))
                    elif values:
                        lines.append(f"{name} — " + ", ".join(str(v) for v in values))
            return normalize_text("\n".join(lines))
        except Exception:
            return ""

    def _ingest_image(self, path: Path, source_id: str, media_type: str) -> IngestionResult:
        warnings: list[str] = []

        # A screenshot/photo of a PDF page arrives with an image extension, but
        # semantically it is still a document. Run a cheap local precheck before
        # asking SigLIP so obvious page scans are not pushed down the generic
        # photograph path. No OCR/model weights are used in this precheck.
        assessment = assess_document_image(path)
        if assessment.is_document_like:
            predictions = [{"label": "document page", "score": assessment.score}]
            routing_source = "document_precheck"
        else:
            if self.siglip is None:
                predictions = self.vlm.classify_file(path, list(DEFAULT_VISUAL_LABELS))
                routing_source = "vlm_router"
            else:
                routing_source = "siglip"
                try:
                    predictions = self.siglip.classify(path)
                except Exception as exc:
                    # SigLIP is an optimization only. A routing failure must not
                    # poison ingestion; use the already-required multimodal VLM.
                    warnings.append(f"SigLIP routing unavailable; used VLM fallback: {exc}")
                    predictions = self.vlm.classify_file(path, list(DEFAULT_VISUAL_LABELS))
                    routing_source = "vlm_fallback"

        top_label = str(predictions[0]["label"]).lower() if predictions else "other visual"

        # If SigLIP is uncertain and a document label is close to the winning
        # score, bias toward document extraction. Missing a document page is much
        # more damaging than giving a natural image a transcription-style prompt.
        if top_label not in self.document_like_image_labels and predictions:
            top_score = float(predictions[0].get("score", 0.0))
            document_prediction = next(
                (p for p in predictions if str(p.get("label", "")).lower() in self.document_like_image_labels),
                None,
            )
            if document_prediction is not None:
                document_score = float(document_prediction.get("score", 0.0))
                if document_score >= 0.15 and document_score >= top_score * 0.85:
                    top_label = str(document_prediction["label"]).lower()
                    routing_source = f"{routing_source}+document_margin"

        if top_label in self.document_like_image_labels:
            prompt = (
                "This image is a document page or screenshot. Transcribe it faithfully for retrieval. "
                "Preserve headings, paragraphs, lists, tables, labels, numbers, equations and reading order. "
                "Describe charts/figures only when they carry information. Do not summarize or invent content."
            )
            strategy = f"image-document:{top_label}"
            kind = "document_image"
        else:
            prompt = None
            strategy = f"image-visual:{top_label}"
            kind = "visual_description"

        description = normalize_text(self.vlm.describe_file(path, prompt=prompt))
        if not description:
            raise RuntimeError("VLM returned an empty description for image input")
        return IngestionResult(
            source_id,
            path.name,
            media_type,
            strategy,
            [SourceElement("visual-0", kind, description, metadata={"visual_route_label": top_label, "processor": "vlm"})],
            warnings=warnings,
            provider_metadata={
                "visual_routing": predictions,
                "routing_source": routing_source,
                "document_precheck": {
                    "is_document_like": assessment.is_document_like,
                    "score": assessment.score,
                    "reasons": list(assessment.reasons),
                },
                "vlm_model_used": getattr(self.vlm, "last_model_used", None),
            },
        )
